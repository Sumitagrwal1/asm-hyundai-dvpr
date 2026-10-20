import openpyxl, glob, os

files = glob.glob(os.path.join(os.path.dirname(os.path.abspath(__file__)), "*.xlsx"))
if files:
    wb = openpyxl.load_workbook(files[0], data_only=True)
    ws = wb["DVPR"] if "DVPR" in wb.sheetnames else wb.active
    for r in list(ws.iter_rows(min_row=3, max_row=8, values_only=True)):
        print(r[1], "|", r[3], "|", r[5], "|", r[15], "| Age:", r[19])
