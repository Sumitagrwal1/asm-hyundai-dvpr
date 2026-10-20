import openpyxl, glob, os, json, re

folder = os.path.dirname(os.path.abspath(__file__))
files = glob.glob(os.path.join(folder, "*.xlsx"))
wb = openpyxl.load_workbook(files[0], data_only=True)
ws = wb["DVPR"]

for r in list(ws.iter_rows(min_row=3, max_row=7, values_only=True)):
    print("EXCEL:", r[1], "|", r[3], "|", r[5], "| Age:", r[19])

with open(os.path.join(folder, "DVPR_Live_Dashboard.html"), encoding="utf-8") as f:
    match = re.search(r"let RAW_DATA = (\[.*?\]);", f.read(), re.DOTALL)
data = json.loads(match.group(1))

for row in data[:4]:
    print("HTML :", row["customer"], "|", row["reg"], "|", row["advisor"], "| Age:", row["age"])

print("Excel rows:", ws.max_row - 2)
print("HTML rows :", len(data))
print("Match:", len(data) == (ws.max_row - 2))
